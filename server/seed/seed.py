"""Build data/owwh.db from salesforce_synthetic_cut.xlsx.

1. Load the 8 real Salesforce Opportunity rows (all 284 columns, original names).
2. Generate synthetic opportunities with the same columns and coherent values.
3. Derive OWWH content (project profiles, case studies, bid submissions, …), contacts,
   sectors, personas and the search index.

By default only the real export is loaded (8 opportunities and the people named in them).
Pass --synthetic to add generated opportunities and derived content for a fuller demo.

Run: .venv/bin/python -m server.seed.seed [--synthetic]
     npm run seed            / npm run seed --synthetic   / npm run seed:synthetic
"""
import argparse
import datetime as dt
import json
import math
import os
import random
import re
import sqlite3
import sys
import uuid
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

from server.app.indexing import index_contact, index_content, index_opportunity
from server.seed import catalog as C
from server.seed.field_meta import DATE_FIELDS, DATETIME_FIELDS, PEOPLE_FIELDS, build_field_meta

ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "salesforce_synthetic_cut.xlsx"
DB_PATH = Path(os.environ.get("OWWH_DB", ROOT / "data" / "owwh.db"))
SCHEMA = ROOT / "server" / "schema.sql"
SYNTHETIC_COUNT = 142
TODAY = dt.date.today()

rng = random.Random(20261019)

CLIENT_SHORT = {
    "Queensland Department of Transport and Main Roads": "TMR", "Transport for NSW": "TfNSW",
    "Main Roads Western Australia": "Main Roads WA", "Department of Transport and Planning Victoria": "DTP",
    "NSW Health Infrastructure": "Health Infrastructure", "Department of Health Victoria": "DH Victoria",
    "Department of Justice and Community Safety": "DJCS", "Victorian Department of Education": "DE Victoria",
    "NSW Department of Education": "NSW Education", "Department of Corrections NZ": "Ara Poutama",
    "Waka Kotahi NZ Transport Agency": "Waka Kotahi", "NBN Co Limited": "NBN", "Optus Networks": "Optus",
    "Department of Industry, Tourism and Trade": "DITT", "NSW Environment Protection Authority": "NSW EPA",
    "Department of Environment and Science QLD": "DES", "Rio Tinto Iron Ore": "Rio Tinto",
    "Department of Finance": "Finance", "Department of Defence": "Defence", "Te Whatu Ora": "Te Whatu Ora",
}

LOCATION_ALIASES = {
    "queensland": "QLD", "qld": "QLD", "victoria": "VIC", "vic": "VIC", "new south wales": "NSW", "nsw": "NSW",
    "western australia": "WA", "wa": "WA", "south australia": "SA", "sa": "SA", "tasmania": "TAS", "tas": "TAS",
    "northern territory": "NT", "nt": "NT", "act": "ACT", "australian capital territory": "ACT",
    "new zealand - north": "NZ North", "new zealand - south": "NZ South",
}


# ---------------------------------------------------------------- helpers
def iso(d):
    if d is None:
        return None
    if isinstance(d, dt.datetime):
        return d.replace(microsecond=0).isoformat()
    return d.isoformat()


def to_date(s):
    if not s:
        return None
    return dt.date.fromisoformat(str(s)[:10])


def add_months(d: dt.date, months: int) -> dt.date:
    y, m = divmod(d.month - 1 + months, 12)
    return dt.date(d.year + y, m + 1, min(d.day, 28))


def client_short(name):
    return CLIENT_SHORT.get(name, name)


def norm_person(name: str) -> str:
    name = name.strip()
    if "," in name:
        last, first = [p.strip() for p in name.split(",", 1)]
        name = f"{first} {last}"
    return name


def slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def money_round(v):
    if v >= 1e7:
        return round(v, -5)
    if v >= 1e6:
        return round(v, -4)
    return round(v, -3)


def sf_id(prefix):
    return prefix + "".join(rng.choice("0123456789") for _ in range(13))


def location_to_regions(loc):
    if not loc:
        return []
    l = loc.lower().strip()
    if l in LOCATION_ALIASES:
        return [LOCATION_ALIASES[l]]
    if "national" in l or l == "australia":
        return ["National"]
    if "new zealand" in l:
        return ["NZ North"]
    return []


# ---------------------------------------------------------------- xlsx
def load_xlsx():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    header = list(rows[0])
    records = []
    for raw in rows[1:]:
        if not any(v is not None for v in raw):
            continue
        rec = {}
        for col, val in zip(header, raw):
            if isinstance(val, bool):
                val = int(val)
            elif isinstance(val, dt.datetime):
                val = iso(val.date()) if col in DATE_FIELDS else iso(val)
            elif isinstance(val, (int, float)) and col in DATE_FIELDS | DATETIME_FIELDS:
                base = dt.datetime(1899, 12, 30) + dt.timedelta(days=float(val))  # Excel serial fallback
                val = iso(base.date()) if col in DATE_FIELDS else iso(base)
            rec[col] = val
        records.append(rec)
    return header, records


def column_types(header, records, meta):
    types = {}
    by_col = {m["column_name"]: m for m in meta}
    for col in header:
        ftype = by_col[col]["field_type"]
        vals = [r[col] for r in records if r.get(col) is not None]
        if ftype in ("date", "datetime", "person", "link", "multiselect"):
            types[col] = "TEXT"
        elif ftype == "boolean":
            types[col] = "INTEGER"
        elif ftype in ("currency", "percent"):
            types[col] = "REAL"
        elif vals and all(isinstance(v, (int, float)) for v in vals):
            types[col] = "REAL"
        else:
            types[col] = "TEXT"
    return types


EMPTY_VALUES = (None, 0, "", "[]", "#Collection(String)")


def used_columns(header, records):
    """Columns with at least one real value in the Salesforce export. Only these are generated."""
    return {c for c in header if any(r.get(c) not in EMPTY_VALUES for r in records)}


def sample_defaults(header, records):
    """Most common value per column in the real rows (e.g. 'No', 0, '[]', None)."""
    return {c: Counter(json.dumps(r.get(c)) for r in records).most_common(1)[0][0] for c in header}


# ---------------------------------------------------------------- synthetic opportunities
SECTOR_WEIGHTS = {"Transport": 28, "Infrastructure Services": 28, "Telecommunications": 20,
                  "Defence & Social Infrastructure": 17, "Corporate Services": 7}
SECTOR_BY_NAME = {s["name"]: s for s in C.SECTORS}


def people_for(sector, titles, exclude=frozenset()):
    """Synthetic pursuit team members. People named in the real export are never reused on synthetic rows."""
    people = [p for p in C.PEOPLE if p[0] not in exclude]
    pool = [p[0] for p in people if (p[1] == sector or p[3] == sector) and any(t in p[2] for t in titles)]
    return pool or [p[0] for p in people if any(t in p[2] for t in titles)]


def fiscal(d: dt.date):
    fy = d.year if d.month >= 7 else d.year - 1
    q = ((d.month - 7) % 12) // 3 + 1
    return fy, q


def fill_template(tpl, ctx):
    # Collapse repeats such as "Yarra Valley Water Water Treatment" when a client name meets the template.
    return re.sub(r"\b(\w+) \1\b", r"\1", tpl.format(**ctx))


def gen_opportunity(header, defaults, used, real_people, brs_counter, uid_counter):
    sector_name = rng.choices(list(SECTOR_WEIGHTS), weights=list(SECTOR_WEIGHTS.values()))[0]
    sector = SECTOR_BY_NAME[sector_name]
    off = rng.choice(sector["offerings"])
    client, client_regions, is_gov = rng.choice(off["clients"])
    region = rng.choice(client_regions)
    months = rng.choice([12, 24, 36, 36, 48, 60, 60, 84, 120])
    ctx = {
        "region": "National" if region == "National" else region,
        "region_name": C.REGION_NAMES[region], "city": C.CITY[region], "client": client,
        "client_short": client_short(client), "highway": rng.choice(C.HIGHWAYS[region]),
        "n": rng.randint(3, 40), "months": months,
    }
    name = fill_template(rng.choice(off["names"]), ctx)
    renewal = rng.random() < 0.35
    if renewal and rng.random() < 0.5:
        name += rng.choice([" - Renewal", " - Extension Option", " (Re-tender)"])
    scope = fill_template(rng.choice(off["scope"]), ctx)

    created = TODAY - dt.timedelta(days=rng.randint(15, 2400))
    age = (TODAY - created).days
    if age > 420:
        outcome = rng.choices(["won", "lost", "open"], weights=[55, 27, 18])[0]
    else:
        outcome = rng.choices(["won", "lost", "open"], weights=[12, 8, 80])[0]

    proposal_close = created + dt.timedelta(days=rng.randint(35, 160))
    if outcome == "open":
        close = max(proposal_close, TODAY) + dt.timedelta(days=rng.randint(20, 240))
        if proposal_close < TODAY and rng.random() < 0.5:
            proposal_close = TODAY + dt.timedelta(days=rng.randint(10, 90))
    else:
        close = min(proposal_close + dt.timedelta(days=rng.randint(20, 140)), TODAY - dt.timedelta(days=1))
        if close < proposal_close:
            proposal_close = close - dt.timedelta(days=14)
    contract_start = close + dt.timedelta(days=rng.randint(0, 60))
    contract_end = add_months(contract_start, months)

    amount = math.exp(rng.gauss(math.log(14e6 if sector_name != "Corporate Services" else 2.5e6), 1.25))
    if "NBN" in client or "Defence" in client:
        amount *= 2.5
    amount = money_round(min(max(amount, 250_000), 420_000_000))

    if outcome == "won":
        stage, prob, fc, fcn, status = "Closed Won", 100, "ClosedWon", "Closed", "Won"
    elif outcome == "lost":
        stage, prob, fc, fcn, status = "Closed Lost", 0, "Omitted", "Omitted", "Lost"
    else:
        stage, prob, fc, fcn = rng.choices(C.STAGES_OPEN, weights=[30, 20, 35, 15])[0]
        status = C.OPEN_STATUS[stage]

    weighted = round(amount * prob / 100, 2)
    acv = money_round(amount / max(1, months / 12))
    restricted = rng.random() < (0.35 if sector_name == "Defence & Social Infrastructure" else 0.12)
    nz = "NZ" in region
    location = {"National": "Australia - National", "NZ North": "New Zealand - North", "NZ South": "New Zealand - South"}.get(region, region)

    o = {col: json.loads(defaults[col]) for col in header}
    brs = next(brs_counter)
    pm = rng.choice(people_for(sector_name, ["Proposal Manager", "Bid Manager"], real_people))
    o.update({
        "@odata.etag": None,
        "ItemInternalId": str(uuid.UUID(int=rng.getrandbits(128), version=4)),
        "Id": sf_id("006Oa"),
        "AccountId": sf_id("001Oa"),
        "OwnerId": sf_id("005Oa"),
        "CreatedById": sf_id("005Oa"),
        "LastModifiedById": "005Oa000006OUhRIAW",
        "RecordTypeId": "012280000005Qn6AAE",
        "Name": name,
        "Short_Opportunity_Name__c": name[:40],
        "Amount": amount,
        "TCV_Converted__c": amount,
        "Estimated_Contract_Value_Rev__c": amount,
        "Annual_Contract_Value__c": acv,
        "ACV_Converted__c": acv,
        "Annual_Weighted_Contract_Value__c": round(acv * prob / 100),
        "Total_Weighted_Contract_Value__c": weighted,
        "CloseDate": iso(close),
        "CreatedDate": iso(dt.datetime.combine(created, dt.time(rng.randint(8, 17), rng.choice([0, 15, 30, 45])))),
        "LastModifiedDate": iso(dt.datetime.combine(min(TODAY, close + dt.timedelta(days=rng.randint(0, 30))), dt.time(9, 30))),
        "CurrencyIsoCode": "NZD" if nz else "AUD",
        "Currency__c": "NZD" if nz else "AUD",
        "Country__c": "New Zealand" if nz else "Australia",
        "StageName": stage,
        "Probability": prob,
        "Go_Probability__c": prob,
        "Get_Probability__c": prob,
        "Go_Get__c": "Go" if prob >= 50 else "Get",
        "ForecastCategory": fc,
        "ForecastCategoryName": fcn,
        "IsClosed": 0 if outcome == "open" else 1,
        "IsWon": 1 if outcome == "won" else 0,
        "Binary_Win__c": 1 if outcome == "won" else 0,
        "Opportunity_Status__c": status,
        "Opportunity_Type__c": "Panel Opportunity" if rng.random() < 0.1 else "Standard Opportunity",
        "Division__c": sector_name,
        "Business_Unit__c": off["bu"],
        "Line_of_Business__c": off["lob"],
        "Discipline__c": off["discipline"],
        "Opportunity_Parent_Account_Name__c": client,
        "Scope_of_Works__c": scope,
        "New_or_Renewal_Work__c": "Renewal Work" if renewal else "New Work",
        "Prospect_Source__c": rng.choice(C.PROSPECT_SOURCES),
        "Proposal_Manager__c": pm,
        "Proposal_Teamsite__c": f"/sites/proposal/{brs:06d}",
        "BRS_Opportunity_Number__c": brs,
        "Unique_Opportunity_Identifier__c": next(uid_counter),
        "Contract_Period_monthsRev1__c": months,
        "Contract_Start_Date__c": iso(contract_start),
        "Contract_End_Date_formula__c": iso(contract_end),
        "Proposal_Closing_Date__c": iso(proposal_close),
        "Tender_Closing_Date__c": iso(proposal_close) if rng.random() < 0.8 else None,
        "Initial_Tender_Close_Date__c": iso(proposal_close - dt.timedelta(days=rng.choice([7, 14, 21]))) if rng.random() < 0.3 else None,
        "Win_Loss_Month__c": iso(close),
        "Opportunity_Age__c": ((close if outcome != "open" else TODAY) - created).days,
        "Project_Location__c": location,
        "Location_of_Project__c": location,
        "Visibility_Classification__c": "Restricted" if restricted else "Unrestricted",
        "Ventia_Legal_Entity__c": "Ventia New Zealand Limited" if nz else "Ventia Australia Pty Ltd",
        "Include_In_New_Work_Forecast__c": "Yes" if outcome == "open" else "No",
        "Binding_Proposal_Required__c": "Yes" if rng.random() < 0.85 else "No",
        "DoA_Level__c": "Level 1" if amount < 2e7 else "Level 2" if amount < 1e8 else "Level 3",
        "PPP_Project__c": "Yes" if "PPP" in name else "No",
        "Panel_Opportunity__c": "Yes" if rng.random() < 0.1 else "No",
        "Parent_Company_Guarantee__c": "Yes" if amount > 1e8 and rng.random() < 0.5 else "No",
        "Fiscal": None,
        "SystemModstamp": None,
    })
    o["SystemModstamp"] = o["LastModifiedDate"]
    fy, q = fiscal(close)
    o.update({"FiscalYear": fy, "FiscalQuarter": q, "Fiscal": f"{fy} {q}"})

    # People (only the people columns the export actually populates)
    if rng.random() < 0.6:
        o["Bid_Writer__c"] = rng.choice(people_for(sector_name, ["Bid Writer", "Proposal Writer"], real_people))
    if rng.random() < 0.45:
        o["Tender_Manager__c"] = rng.choice(people_for(sector_name, ["Tender Manager", "Bid Manager"], real_people))

    # Outcome detail
    if outcome != "open":
        reasons = rng.sample(C.WIN_REASONS if outcome == "won" else C.LOSS_REASONS, rng.randint(1, 2))
        if outcome == "won" and renewal and "Incumbent" not in reasons:
            reasons[0] = "Incumbent"
        o["Won_Lost_Reason__c"] = json.dumps(reasons)
        o["Reason_Won_Lost__c"] = reasons[0]
    o["Ventia_JV_Portion__c"] = 100

    # Region revenue split: the export only populates the QLD and VIC split columns, so only those are written.
    prefix = C.REGION_PREFIX.get(region)
    if prefix and f"{prefix}__c" in used:
        o[f"{prefix}__c"] = 100
        o[f"{prefix}_Revenue__c"] = amount
        o[f"{prefix}_Weighted_Revenue__c"] = round(amount * prob / 100, 2)

    # Hard rule: columns that are empty across the whole export stay exactly as the export has them.
    for col in header:
        if col not in used:
            o[col] = json.loads(defaults[col])
    return o


def facets_for(o, source):
    regions = {}
    for r, prefix in C.REGION_PREFIX.items():
        pct = o.get(f"{prefix}__c") or 0
        if pct:
            regions[r] = pct
    if not regions:
        regions = {r: 100 for r in location_to_regions(o.get("Project_Location__c") or o.get("Location_of_Project__c"))}
    if o.get("Project_Location__c") == "Australia - National" or len(regions) >= 3:
        primary = "National"
    else:
        primary = max(regions, key=regions.get) if regions else None
    status = "won" if o.get("IsWon") else ("lost" if o.get("IsClosed") else "open")
    tcv = o.get("TCV_Converted__c") or o.get("Amount") or o.get("Estimated_Contract_Value_Rev__c") or None
    return {
        "opportunity_id": o["ItemInternalId"],
        "sector": o.get("Division__c"),
        "client": o.get("Opportunity_Parent_Account_Name__c"),
        "region": primary,
        "regions": ",".join(regions),
        "service_line": o.get("Line_of_Business__c") or o.get("Discipline__c") or o.get("Business_Unit__c"),
        "status": status,
        "tcv": tcv,
        "is_restricted": 1 if o.get("Visibility_Classification__c") == "Restricted" else 0,
        "source": source,
    }


# ---------------------------------------------------------------- contacts
ROLE_EXPERTISE = {
    "Proposal Manager": ["Work Winning"], "Bid Writer": ["Work Winning"], "Tender Manager": ["Work Winning"],
    "Tender Co-ordinator": ["Work Winning"], "Lead Estimator": ["Commercial", "Estimating"],
    "Commercial Lead": ["Commercial"], "Legal Lead": ["Legal"], "Delivery Lead": ["Operational", "Technical"],
    "Technical Lead": ["Technical"],
}

DELIVERY_PEOPLE = {
    "Transport": [("Grant Hollis", "Project Director - Roads"), ("Tama Parata", "Rail Operations Manager")],
    "Infrastructure Services": [("Dr. Anika Rao", "Principal Process Engineer - Water"), ("Steve Mullins", "Operations Manager - Resources")],
    "Telecommunications": [("Kenji Watanabe", "Network Engineering Lead"), ("Bree Holloway", "Field Operations Manager")],
    "Defence & Social Infrastructure": [("Col. (Ret.) Mark Ellison", "Estate Services Director"), ("Rebecca Stone", "Hospital FM Operations Manager")],
    "Corporate Services": [("Arjun Mehta", "Head of Digital Asset Intelligence"), ("Lucy Tran", "Sustainability Advisory Lead")],
}


def build_contacts(opps, facets_by_id, real_people):
    # Catalog titles only apply to synthetic people; people from the export get the role they hold in the data.
    title_by_name = {p[0]: p[2] for p in C.PEOPLE if p[0] not in real_people}
    data = defaultdict(lambda: {"roles": Counter(), "sectors": Counter(), "clients": Counter(), "service_lines": Counter(),
                                "regions": Counter(), "opps": set(), "wins": set(), "closed": set(), "tcv": 0.0, "source": None})
    links = []
    for o in opps:
        f = facets_by_id[o["ItemInternalId"]]
        for col, role in PEOPLE_FIELDS.items():
            raw = o.get(col)
            if not raw or not isinstance(raw, str):
                continue
            name = norm_person(raw)
            d = data[name]
            d["source"] = d["source"] or raw
            d["roles"][role] += 1
            if f["sector"]:
                d["sectors"][f["sector"]] += 1
            if f["client"]:
                d["clients"][f["client"]] += 1
            if f["service_line"]:
                d["service_lines"][f["service_line"]] += 1
            if f["region"]:
                d["regions"][f["region"]] += 1
            if o["ItemInternalId"] not in d["opps"]:
                d["opps"].add(o["ItemInternalId"])
                d["tcv"] += f["tcv"] or 0
                if f["status"] == "won":
                    d["wins"].add(o["ItemInternalId"])
                if f["status"] != "open":
                    d["closed"].add(o["ItemInternalId"])
            links.append((slug(name), o["ItemInternalId"], role))

    contacts = []
    for name, d in data.items():
        primary_sector = d["sectors"].most_common(1)[0][0] if d["sectors"] else None
        top_role = d["roles"].most_common(1)[0][0]
        expertise = sorted({e for r in d["roles"] for e in ROLE_EXPERTISE.get(r, [])})
        title = title_by_name.get(name) or ("Tender Collateral Team" if name == C.PROPOSAL_TEAM else top_role)
        top_clients = [c for c, _ in d["clients"].most_common(5)]
        region = d["regions"].most_common(1)[0][0] if d["regions"] else "National"
        bio = (f"{title} in {primary_sector or 'the business'} with experience across "
               f"{len(d['opps'])} pursuits"
               + (f", including work for {', '.join(top_clients[:3])}." if top_clients else "."))
        contacts.append({
            "id": slug(name), "name": name, "source_name": d["source"], "title": title,
            "location": C.REGION_NAMES.get(region, region), "primary_sector": primary_sector,
            "sectors": json.dumps([s for s, _ in d["sectors"].most_common()]),
            "expertise": json.dumps(expertise), "roles": json.dumps([r for r, _ in d["roles"].most_common()]),
            "clients": json.dumps(top_clients), "service_lines": json.dumps([s for s, _ in d["service_lines"].most_common(5)]),
            "bio": bio, "opp_count": len(d["opps"]), "win_count": len(d["wins"]), "closed_count": len(d["closed"]),
            "total_tcv": d["tcv"], "is_team": 1 if "Team" in name else 0,
        })
    return contacts, links


def add_delivery_people(contacts, links, won_by_sector):
    for sector, people in DELIVERY_PEOPLE.items():
        for name, title in people:
            role = "Delivery Lead" if "Manager" in title or "Director" in title else "Technical Lead"
            opps = rng.sample(won_by_sector.get(sector, []), min(6, len(won_by_sector.get(sector, []))))
            for o in opps:
                links.append((slug(name), o["ItemInternalId"], role))
            contacts.append({
                "id": slug(name), "name": name, "source_name": name, "title": title,
                "location": rng.choice(["Brisbane", "Sydney", "Melbourne", "Perth", "Auckland"]),
                "primary_sector": sector, "sectors": json.dumps([sector]),
                "expertise": json.dumps(ROLE_EXPERTISE[role] + (["Sector Specialist"])),
                "roles": json.dumps([role]),
                "clients": json.dumps(list(dict.fromkeys(o["Opportunity_Parent_Account_Name__c"] for o in opps))[:5]),
                "service_lines": json.dumps(list(dict.fromkeys(o["Line_of_Business__c"] for o in opps if o.get("Line_of_Business__c")))[:5]),
                "bio": f"{title} who has led delivery of {len(opps)} contracts in {sector}. Go-to contact for technical and operational input to bids.",
                "opp_count": len(opps), "win_count": len(opps), "closed_count": len(opps),
                "total_tcv": sum(o.get("TCV_Converted__c") or 0 for o in opps), "is_team": 0,
                "data_source": "synthetic",
            })


# ---------------------------------------------------------------- content
OUTCOME_LINES = [
    "Delivered {pct}% of planned works on time across the first contract year.",
    "Zero lost-time injuries over {n} million hours worked.",
    "Achieved {pct}% client satisfaction score in the annual performance review.",
    "Reduced reactive maintenance call-outs by {pct2}% through predictive maintenance.",
    "Engaged {n2} local and Indigenous subcontractors, exceeding social procurement targets.",
    "Delivered {pct2}% cost saving to the client through innovation in delivery methodology.",
    "Contract extended by the client following strong KPI performance.",
]
METHOD_STEPS = [
    "Mobilisation plan with a 30/60/90 day transition and knowledge transfer from the incumbent.",
    "Risk-based asset maintenance strategy driven by condition data and criticality scoring.",
    "Integrated safety management system aligned to ISO 45001 with daily pre-start briefings.",
    "Digital work management with mobile field apps and real-time client reporting dashboards.",
    "Local workforce and supply chain plan including Indigenous participation targets.",
    "Continuous improvement framework with quarterly innovation workshops with the client.",
    "Environmental management plan covering waste, water quality and community impacts.",
]
LESSONS_WON = [
    "Early client engagement before RFT release shaped a compliant, differentiated solution.",
    "Case studies with quantified outcomes scored strongly in the evaluation.",
    "Bringing the proposed delivery lead to the interview built client confidence.",
    "A clear transition plan reduced the perceived risk of changing contractor.",
]
LESSONS_LOST = [
    "Pricing was above the client's pre-tender estimate; earlier cost validation needed.",
    "Technical response lacked local project examples; reuse more sector case studies.",
    "Incumbent relationship was underestimated; start capture planning earlier.",
    "Response was too generic - tailor methodology to the client's stated priorities.",
]
CLIENT_QUESTIONS = [
    ("Describe your approach to safety leadership on this contract.",
     "Our safety approach is led from the top with visible leadership, critical-risk controls and daily engagement with frontline crews."),
    ("How will you ensure a smooth transition from the incumbent?",
     "We will run a structured 90-day transition with dedicated transition manager, early workforce engagement and asset data validation."),
    ("Outline how you will deliver value for money over the term.",
     "We will use condition-based planning, bundled works and an open-book innovation register to share savings with the client."),
    ("Describe your local industry and social procurement commitments.",
     "We commit to local subcontracting targets, Indigenous business participation and apprenticeship pathways in the region."),
]

TEMPLATES = [
    ("Executive Summary Template", "Standard structure for tender executive summaries: client challenge, our solution, why us, value."),
    ("Safety Management Response Template", "Pre-approved safety leadership and critical risk management response with sector variants."),
    ("Transition-In Plan Template", "90-day transition plan template including RACI, risk register and communication plan."),
    ("Social Procurement Response Template", "Response template for Indigenous participation, local content and social value questions."),
]
GOOD_NEWS = [
    ("{client_short} Recognises Team for Outstanding Safety Performance", "The {bu} team was recognised by {client} for 1,000 days without a lost-time injury."),
    ("Apprentice Program Expands in {city}", "Our {bu} team welcomed {n} new apprentices in {city}, supporting local employment."),
    ("Innovation Award for Digital Maintenance at {client_short}", "Our digital maintenance platform at {client_short} won an industry innovation award."),
]


def make_dates(content_type):
    cycle = C.CONTENT_REVIEW_MONTHS[content_type]
    offset = rng.randint(-75, cycle * 30 - 10)
    review = TODAY + dt.timedelta(days=offset)
    updated = add_months(review, -cycle)
    expiry = add_months(review, 12)
    if rng.random() < 0.04:
        expiry = TODAY - dt.timedelta(days=rng.randint(5, 60))
    return review, updated, expiry


def choose_visibility(sector, flags):
    if flags:
        return "sector_only", []
    r = rng.random()
    if r < 0.45:
        return "enterprise", []
    if r < 0.72:
        others = [s["name"] for s in C.SECTORS if s["name"] != sector]
        return "shared", rng.sample(others, rng.randint(1, 2))
    return "sector_only", []


def content_from_opp(o, f, ctype, sector_meta):
    tcv = f["tcv"]
    region_name = C.REGION_NAMES.get(f["region"] or "National", f["region"])
    client = f["client"] or "Client"
    flags = []
    if f["is_restricted"]:
        flags += ["client_sensitive", "nda"]
    if ctype == "bid_submission" and rng.random() < 0.25:
        flags.append("pricing")
    if ctype == "bid_submission" and f["status"] == "open":
        flags.append("current_tender")
    gov_client = any(k in client for k in ("Department", "Defence", "Council", "Health", "Transport", "Water", "Waka"))
    if gov_client and rng.random() < 0.12:
        flags.append("government")
    if ctype == "case_study" and rng.random() < 0.12:
        flags.append("referee")
    flags = sorted(set(flags))
    visibility, shares = choose_visibility(f["sector"], flags)
    review, updated, expiry = make_dates(ctype)
    months = o.get("Contract_Period_monthsRev1__c")
    scope = o.get("Scope_of_Works__c") or f"{o.get('Discipline__c') or f['service_line'] or 'Services'} delivered for {client}."
    fmt = {"pct": rng.randint(92, 99), "pct2": rng.randint(8, 25), "n": rng.randint(1, 6), "n2": rng.randint(5, 40)}
    reasons = json.loads(o.get("Won_Lost_Reason__c") or "[]")
    services = [s for s in [o.get("Discipline__c"), o.get("Line_of_Business__c"), o.get("Business_Unit__c")] if s]
    pm = o.get("Proposal_Manager__c")
    tags = [t for t in [o.get("Discipline__c"), o.get("New_or_Renewal_Work__c"), f["region"]] if t]

    if ctype == "project_profile":
        title = o["Name"]
        summary = f"{scope}"
        body = {
            "description": scope, "services_delivered": services, "duration_months": months,
            "start_date": o.get("Contract_Start_Date__c"), "end_date": o.get("Contract_End_Date_formula__c"),
            "contract_value": tcv, "location": region_name, "outcome_status": "won",
        }
        confidentiality = "Internal"
    elif ctype == "case_study":
        title = f"Case Study: {o['Name']}"
        outcomes = [l.format(**fmt) for l in rng.sample(OUTCOME_LINES, 3)]
        summary = f"How we delivered {o.get('Discipline__c') or 'services'} for {client} in {region_name}. {outcomes[0]}"
        body = {
            "challenge": f"{client} needed a partner to deliver {(o.get('Discipline__c') or 'services').lower()} across {region_name} with improved safety, reliability and value for money.",
            "solution": scope, "scope": scope, "service_offering": services[0] if services else None,
            "outcomes": outcomes, "location": region_name, "photo_seed": rng.randint(1, 9999),
            "reference_contact": {"name": rng.choice(["Jane Mitchell", "Peter Grant", "Aisha Karim", "Tom Hughes", "Linda Park"]),
                                  "title": rng.choice(["Contract Manager", "Director of Asset Services", "Program Manager"]),
                                  "organisation": client},
            "contract_value": tcv, "duration_months": months, "outcome_status": "won",
        }
        confidentiality = "Restricted" if flags else "Internal"
    else:  # bid_submission
        won = f["status"] == "won"
        title = f"Bid Submission: {o['Name']}"
        summary = (f"{'Successful' if won else 'Unsuccessful' if f['status'] == 'lost' else 'Current'} submission to {client}. "
                   + (f"Key {'win' if won else 'loss'} factor: {reasons[0]}." if reasons else ""))
        body = {
            "executive_summary": f"We proposed a {months or 36}-month solution to {scope[0].lower() + scope[1:]} Our offer combined proven local delivery, a strong safety culture and a digital-first maintenance approach.",
            "methodology": rng.sample(METHOD_STEPS, 4),
            "client_responses": [{"question": q, "answer": a} for q, a in rng.sample(CLIENT_QUESTIONS, 2)],
            "outcome": "Won" if won else "Lost" if f["status"] == "lost" else "Pending",
            "outcome_status": f["status"], "win_loss_reasons": reasons,
            "lessons_learned": rng.sample(LESSONS_WON if won else LESSONS_LOST, 2),
            "submitted_date": o.get("Proposal_Closing_Date__c"), "tender_number": o.get("Client_s_Tender_Number__c"),
            "contract_value": tcv, "teamsite": o.get("Proposal_Teamsite__c"),
        }
        confidentiality = "Restricted" if flags else "Confidential"

    return {
        "id": f"{ctype[:2]}-{o['ItemInternalId'][:8]}",
        "type": ctype, "title": title, "summary": summary, "body_json": json.dumps(body),
        "sector": f["sector"], "business_unit": o.get("Business_Unit__c") or sector_meta["offerings"][0]["bu"],
        "service_line": f["service_line"], "contract_value": tcv, "contract_value_range": C.value_range(tcv),
        "client": client, "geography": f["region"], "confidentiality": confidentiality,
        "visibility": visibility, "restricted_flags": json.dumps(flags), "is_restricted": 1 if flags else 0,
        "owner": norm_person(pm) if pm else sector_meta["content_owner"],
        "review_date": iso(review), "expiry_date": iso(expiry), "approval_status": "approved",
        "version": rng.randint(1, 4), "source_opportunity_id": o["ItemInternalId"], "tags": json.dumps(tags),
        "created_by": None, "created_at": iso(dt.datetime.combine(to_date(o.get("CloseDate")) or TODAY, dt.time(10))),
        "updated_at": iso(dt.datetime.combine(min(updated, TODAY), dt.time(10))),
        "_shares": shares,
    }


def sector_content(sector, sector_meta, opps_in_sector):
    items = []
    off = sector_meta["offerings"][0]
    client = off["clients"][0][0]
    region = off["clients"][0][1][0]
    ctx = {"client": client, "client_short": client_short(client), "bu": off["bu"], "city": C.CITY[region], "n": rng.randint(6, 25)}

    def base(ctype, title, summary, body, confidentiality, visibility="enterprise", shares=None):
        review, updated, expiry = make_dates(ctype)
        return {
            "id": f"{ctype[:2]}-{sector_meta['id']}-{slug(title)[:24]}",
            "type": ctype, "title": title, "summary": summary, "body_json": json.dumps(body), "sector": sector,
            "business_unit": off["bu"], "service_line": off["lob"], "contract_value": None, "contract_value_range": "Not applicable",
            "client": "Various clients" if ctype in ("template", "capability_statement") else client, "geography": region,
            "confidentiality": confidentiality, "visibility": visibility, "restricted_flags": "[]", "is_restricted": 0,
            "owner": sector_meta["content_owner"], "review_date": iso(review), "expiry_date": iso(expiry),
            "approval_status": "approved", "version": rng.randint(1, 5), "source_opportunity_id": None,
            "tags": json.dumps([off["discipline"]]), "created_by": None,
            "created_at": iso(dt.datetime.combine(add_months(updated, -rng.randint(0, 12)), dt.time(9))),
            "updated_at": iso(dt.datetime.combine(min(updated, TODAY), dt.time(9))), "_shares": shares or [],
        }

    for title, desc in rng.sample(TEMPLATES, 2):
        items.append(base("template", f"{sector_meta['short']} {title}", desc,
                          {"description": desc, "sections": ["Context", "Our approach", "Evidence", "Commitments"]}, "Internal",
                          visibility=rng.choice(["enterprise", "sector_only"])))
    for title, story in rng.sample(GOOD_NEWS, 2):
        items.append(base("good_news_story", title.format(**ctx), story.format(**ctx),
                          {"story": story.format(**ctx), "highlights": [f"{ctx['n']} team members recognised", f"Client: {client}"]}, "Public"))
    caps = sorted({o.get("Discipline__c") for o in opps_in_sector if o.get("Discipline__c")})[:6]
    items.append(base("capability_statement", f"{sector} Capability Statement", sector_meta["description"],
                      {"statement": sector_meta["description"], "capabilities": caps}, "Public"))
    for i in range(2):
        o = rng.choice(opps_in_sector)
        items.append(base("image", f"Site photo: {o['Name'][:60]}", f"Approved project photography from {o['Name']}.",
                          {"caption": f"Crew on site - {o.get('Opportunity_Parent_Account_Name__c')}", "photo_seed": rng.randint(1, 9999),
                           "photographer": "Corporate Communications"}, "Internal",
                          visibility="shared", shares=rng.sample([s["name"] for s in C.SECTORS if s["name"] != sector], 1)))
        items[-1]["source_opportunity_id"] = o["ItemInternalId"]
    return items


# ---------------------------------------------------------------- main
def synthetic_requested() -> bool:
    parser = argparse.ArgumentParser(description="Build data/owwh.db from salesforce_synthetic_cut.xlsx")
    parser.add_argument("--synthetic", action="store_true", help="also generate synthetic opportunities and content")
    args = parser.parse_args()
    # `npm run setup --synthetic` reaches this script as npm_config_synthetic=true (npm turns unknown flags into
    # config env vars, and nested `npm run seed` inherits them).
    return args.synthetic or os.environ.get("npm_config_synthetic", "").lower() in ("true", "1")


def main(synthetic: bool):
    header, real = load_xlsx()
    meta = build_field_meta(header)
    types = column_types(header, real, meta)
    defaults = sample_defaults(header, real)
    used = used_columns(header, real)
    real_people = {norm_person(r[c]) for r in real for c in PEOPLE_FIELDS if isinstance(r.get(c), str)}

    brs_counter = iter(range(16500, 99999, 7))
    uid_counter = iter(range(3000, 99999))
    generated = [gen_opportunity(header, defaults, used, real_people, brs_counter, uid_counter)
                 for _ in range(SYNTHETIC_COUNT)] if synthetic else []
    opps = real + generated

    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    if DB_PATH.exists():
        DB_PATH.unlink()
    db = sqlite3.connect(DB_PATH)
    db.row_factory = sqlite3.Row

    cols_sql = ",\n  ".join(f'"{c}" {types[c]}' + (" PRIMARY KEY" if c == "ItemInternalId" else "") for c in header)
    db.execute(f"CREATE TABLE opportunities (\n  {cols_sql}\n)")
    db.executescript(SCHEMA.read_text())

    db.executemany("INSERT INTO sf_field_meta VALUES (:column_name, :label, :field_group, :field_type, :sort_order, :is_key, :display)", meta)
    placeholders = ", ".join("?" for _ in header)
    quoted = ", ".join(f'"{c}"' for c in header)
    db.executemany(f"INSERT INTO opportunities ({quoted}) VALUES ({placeholders})", [[o.get(c) for c in header] for o in opps])

    real_ids = {o["ItemInternalId"] for o in real}
    facets = [facets_for(o, "salesforce" if o["ItemInternalId"] in real_ids else "synthetic") for o in opps]
    facets_by_id = {f["opportunity_id"]: f for f in facets}
    db.executemany("INSERT INTO opportunity_facets VALUES (:opportunity_id, :sector, :client, :region, :regions, :service_line, :status, :tcv, :is_restricted, :source)", facets)

    # Sectors + personas
    personas_by_sector = defaultdict(list)
    for p in C.PERSONAS:
        personas_by_sector[p[4]].append(p)
    for s in C.SECTORS:
        seniors = [p[0] for p in C.PEOPLE if p[1] == s["name"] and ("Senior" in p[2] or "Bid Manager" in p[2] or "Proposal Manager" in p[2])]
        admin = next(p[1] for p in personas_by_sector[s["name"]] if p[3] == "sector_admin")
        others = [p[0] for p in C.PEOPLE if p[1] == s["name"] and p[0] not in seniors[:1]]
        # Content owner / backup are invented names, so they are only set in synthetic mode.
        s["content_owner"] = seniors[0] if synthetic else None
        db.execute("INSERT INTO sectors VALUES (?, ?, ?, ?, ?, ?, ?)",
                   (s["id"], s["name"], s["short"], s["description"], s["content_owner"], admin, others[0] if synthetic else None))
    db.executemany("INSERT INTO users VALUES (?, ?, ?, ?, ?)", C.PERSONAS)

    # Contacts
    contacts, links = build_contacts(opps, facets_by_id, real_people)
    for c in contacts:
        c["data_source"] = "salesforce" if c["name"] in real_people else "synthetic"
    won_by_sector = defaultdict(list)
    for o in opps:
        if facets_by_id[o["ItemInternalId"]]["status"] == "won":
            won_by_sector[o["Division__c"]].append(o)
    if synthetic:
        add_delivery_people(contacts, links, won_by_sector)
    contact_cols = list(contacts[0])
    db.executemany(f"INSERT INTO contacts ({', '.join(contact_cols)}) VALUES ({', '.join(':' + c for c in contact_cols)})", contacts)
    db.executemany("INSERT OR IGNORE INTO contact_opportunities VALUES (?, ?, ?)", links)
    links_by_opp = defaultdict(list)
    for cid, oid, role in links:
        links_by_opp[oid].append((cid, role))

    # Content: every item is generated text, so none is created from the real export alone.
    content = []
    if synthetic:
        for o in opps:
            f = facets_by_id[o["ItemInternalId"]]
            sm = SECTOR_BY_NAME.get(f["sector"])
            if not sm:
                continue
            if f["status"] == "won":
                content.append(content_from_opp(o, f, "project_profile", sm))
                if (f["tcv"] or 0) > 1e6 and rng.random() < 0.45 or o["ItemInternalId"] in real_ids:
                    content.append(content_from_opp(o, f, "case_study", sm))
            if f["status"] in ("won", "lost") or (f["status"] == "open" and o.get("StageName") == "Negotiation/Review"):
                content.append(content_from_opp(o, f, "bid_submission", sm))
        for s in C.SECTORS:
            content += sector_content(s["name"], s, [o for o in opps if o.get("Division__c") == s["name"]])

        # A few items awaiting approval / rejected so governance has a queue
        by_sector = defaultdict(list)
        for c in content:
            by_sector[c["sector"]].append(c)
        for sector, items in by_sector.items():
            for c in rng.sample(items, min(2, len(items))):
                c["approval_status"] = "pending"
                c["version"] = 1
            if items and rng.random() < 0.6:
                rng.choice([c for c in items if c["approval_status"] == "approved"])["approval_status"] = "rejected"
        # One item per sector with missing metadata to make completeness scores meaningful
        for sector, items in by_sector.items():
            c = rng.choice(items)
            c["business_unit"] = None if rng.random() < 0.5 else c["business_unit"]
            c["geography"] = None if c["business_unit"] else c["geography"]

    ccols = [k for k in content[0] if not k.startswith("_")] if content else []
    for c in content:
        db.execute(f"INSERT INTO content_items ({', '.join(ccols)}) VALUES ({', '.join('?' for _ in ccols)})", [c[k] for k in ccols])
        for s in c["_shares"]:
            db.execute("INSERT INTO content_shares VALUES (?, ?)", (c["id"], s))
        roles_seen = set()
        for cid, role in links_by_opp.get(c["source_opportunity_id"], []):
            if (cid, role) not in roles_seen:
                roles_seen.add((cid, role))
                db.execute("INSERT OR IGNORE INTO content_contacts VALUES (?, ?, ?)", (c["id"], cid, role))
        if not c["source_opportunity_id"] or not roles_seen:
            owner_id = slug(c["owner"]) if c["owner"] else None
            if owner_id and db.execute("SELECT 1 FROM contacts WHERE id = ?", (owner_id,)).fetchone():
                db.execute("INSERT OR IGNORE INTO content_contacts VALUES (?, ?, ?)", (c["id"], owner_id, "Content Owner"))
        for v in range(1, c["version"] + 1):
            when = to_date(c["created_at"]) + dt.timedelta(days=(v - 1) * rng.randint(60, 200))
            db.execute("INSERT INTO content_versions (content_id, version, changed_by, changed_at, summary) VALUES (?, ?, ?, ?, ?)",
                       (c["id"], v, c["owner"], iso(dt.datetime.combine(min(when, TODAY), dt.time(11))),
                        "Initial publication" if v == 1 else rng.choice(["Updated outcomes and metrics", "Annual review - no changes",
                                                                          "Refreshed photos and reference contact", "Metadata corrected"])))

    # Search index
    for c in content:
        index_content(db, c)
    for o in opps:
        index_opportunity(db, o, facets_by_id[o["ItemInternalId"]])
    for c in contacts:
        index_contact(db, c)

    db.execute("INSERT INTO audit_log (ts, user_id, action, entity_type, entity_id, detail) VALUES (?, ?, ?, ?, ?, ?)",
               (iso(dt.datetime.now()), "system", "seed", "database", str(DB_PATH.name),
                json.dumps({"opportunities": len(opps), "real": len(real), "synthetic": synthetic, "content": len(content)})))
    db.commit()

    counts = {t: db.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in
              ["opportunities", "sf_field_meta", "contacts", "content_items", "content_shares", "content_contacts", "search_docs"]}
    by_type = dict(db.execute("SELECT type, COUNT(*) FROM content_items GROUP BY type").fetchall())
    mode = "real + synthetic" if synthetic else "real data only (add --synthetic for demo data)"
    print(f"Seeded {DB_PATH.relative_to(ROOT) if DB_PATH.is_relative_to(ROOT) else DB_PATH} [{mode}]: {len(real)} real + {len(generated)} synthetic opportunities "
          f"({len(header)} Salesforce columns)")
    for k, v in counts.items():
        print(f"  {k:22} {v}")
    print("  content by type:", by_type)
    db.close()


if __name__ == "__main__":
    sys.exit(main(synthetic_requested()))
